import sys
import os
sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..'))
from kite_trade import *
import threading
import time
import datetime
import winsound
import os
import json
from flask import *
import pandas as pd
app = Flask(__name__)

# # Second way is provide 'enctoken' manually from 'kite.zerodha.com' website
# # Than you can use login window of 'kite.zerodha.com' website Just don't logout from that window
# currStk = next((item for item in stocks if item['stock'] == stockName), None)
enctoken = ""

# kite client instances will be created once enctoken is provided
kite = None
kiteNew = None


data1 = {
    "stock": {
        "stockName": "GLENMARK", # Need to Set
        "exchange": "NSE",
        "instrumentToken": "",
        "superTrend": "103",
        "stoploss": 100,
    },
    "common": {
        "dateTime": datetime.datetime.now().strftime("%Y-%m-%d %I:%M:%S %p"),
        "message": "",
        "transactionType": "",
        "availableFund": 0,
        "availableQuantity": 0,
        "plannedQuantity": 1, # Need to Set,
        "currentPnl": 0,
    },
    "trendInfo": {
        "o": 0,
        "h": 0,
        "l": 0,
        "c": 0,
        "a": 0,
        "sl": 0,
        "recentT": [],
        "previousT": "",
        "currentT": "",
    },
    "StopAlgo": False
}

class Equity:
    def setEnctoken(self, encToken):
        # initialize global enctoken and kite clients
        global enctoken, kite, kiteNew
        enctoken = encToken
        kite = KiteApp(enctoken=enctoken)
        kiteNew = ZerodhaKiteTrade()
    def placeOrder(self, transactionType, additionalQty=0):
        if data1["common"]["availableQuantity"] == 0 or data1["common"]["availableQuantity"] == -(data1["common"]["plannedQuantity"]) or data1["common"]["availableQuantity"] == (data1["common"]["plannedQuantity"]):
            order = kite.place_order(variety=kite.VARIETY_REGULAR,
                                exchange=kite.EXCHANGE_NSE,
                                tradingsymbol=data1["stock"]["stockName"],
                                transaction_type=transactionType,#kite.TRANSACTION_TYPE_BUY,
                                quantity=(data1["common"]["plannedQuantity"] + additionalQty),
                                product=kite.PRODUCT_MIS,
                                order_type=kite.ORDER_TYPE_MARKET,
                                price=None,
                                validity=None,
                                disclosed_quantity=None,
                                trigger_price=None,
                                squareoff=None, 
                                stoploss=None,
                                trailing_stoploss=None,
                                tag = "SuperTrendAlgo"
                                )            
            self.currentPostions()
            return order
        else:
            data1["common"]["message"] = "Quantity limit has been Reached Please check"
            return None
    def currentFund(self):                   
        data1["common"]["availableFund"] = kite.margins().get('equity').get('available').get("live_balance") 
    def currentPostions(self):        
        positions = kite.positions()        
        if positions is not None:       
            currStk = next((item for item in positions['day'] if item['tradingsymbol'] == data1["stock"]["stockName"] and item['exchange'] == data1["stock"]["exchange"] and item['product'] == 'MIS'), None)
            if currStk is not None:
                data1["common"]["availableQuantity"] = currStk['quantity']
                data1["common"]["currentPnl"] = currStk['pnl'] 
    def setInstrumentToken(self): 
        try: 
            instrument_token = kite.instruments(data1["stock"]["exchange"], data1["stock"]["stockName"])
            if not instrument_token:
                print("No instruments returned for", data1["stock"]["stockName"])
                data1["common"]["message"] = "Instrument not found"
                return

            # safe extraction
            token0 = instrument_token[0]
            if not isinstance(token0, dict) or "instrument_token" not in token0:
                print("Unexpected instrument format:", token0)
                data1["common"]["message"] = "Instrument format error"
                return

            data1["stock"]["instrumentToken"] = token0.get("instrument_token")

            from_datetime = datetime.datetime.now() - datetime.timedelta(days=5)
            to_datetime = datetime.datetime.now()

            history = kite.historical_data(data1["stock"]["instrumentToken"], from_datetime, to_datetime, "5minute", continuous=False, oi=False)

            # ensure history is a DataFrame
            if isinstance(history, list):
                history = pd.DataFrame(history)

            if history is None or len(history) == 0:
                print("No historical data returned for token", data1["stock"]["instrumentToken"])
                data1["common"]["message"] = "No history"
                return

            required_cols = {'open', 'high', 'low', 'close'}
            if not required_cols.issubset(set(history.columns)):
                print("History missing required columns:", history.columns)
                data1["common"]["message"] = "History columns missing"
                return

            # ensure there is at least one row before using iloc[-1]
            last = history.iloc[-1]
            data1["trendInfo"]["a"] = (last['open'] + last['high'] + last['low'] + last['close']) / 4
            data1["trendInfo"]["o"] = last['open']
            data1["trendInfo"]["h"] = last['high']
            data1["trendInfo"]["l"] = last['low']
            data1["trendInfo"]["c"] = last['close']

            # call SuperTrend only if enough rows for the period
            min_rows = 1  # adjust if SuperTrend requires more (e.g., period)
            if len(history) < min_rows:
                print("Not enough history rows for SuperTrend:", len(history))
                return
            
            length = 10
            factory = 3
            if (data1["stock"]["superTrend"] == "73"):
                length = 7
                factory = 3
                
            trendInfo = kiteNew.SuperTrend(history, length, factory)

            # store last trend dataframe for later decision logic
            self.last_trend_df = trendInfo

            # validate SuperTrend output before indexing
            if 'STX' not in trendInfo.columns:
                print("SuperTrend result missing STX:", trendInfo.columns)
                data1["common"]["message"] = "SuperTrend error"
                return

            if trendInfo["STX"].iloc[-1]:  # Uptrend
                if 'final_lb' in trendInfo.columns:
                    data1["trendInfo"]["sl"] = trendInfo["final_lb"].iloc[-1]
                else:
                    print("final_lb missing in SuperTrend output")
            else:  # Downtrend
                if 'final_ub' in trendInfo.columns:
                    data1["trendInfo"]["sl"] = trendInfo["final_ub"].iloc[-1]
                else:
                    print("final_ub missing in SuperTrend output")

            data1["trendInfo"]["currentT"] = trendInfo["STX"].iloc[-1]

            if len(data1["trendInfo"]["recentT"]) >= 3:
                data1["trendInfo"]["recentT"].pop(0)
                data1["trendInfo"]["recentT"].append(data1["trendInfo"]["currentT"])
            else:
                data1["trendInfo"]["recentT"].append(data1["trendInfo"]["currentT"])

        except Exception as inst:
            print(type(inst))
            print(inst.args)
            print(inst)
    def setAllData(self):
        self.currentFund() 
        self.currentPostions()
        self.setInstrumentToken()

    # new method: decide_action - returns dict or None
    def decide_action(self, atr_multiplier=1.5, max_layers=2, max_risk_qty=None):
        """
        Decide next action based on data1 state and last_trend_df.
        Returns:
            dict: { action: 'BUY'|'SELL', reverse: bool, qty: int, stop: float, reason: str }
            or None when no action
        """
        # require trend dataframe available
        trend_df = getattr(self, "last_trend_df", None)
        if trend_df is None:
            return None

        recent = data1["trendInfo"]["recentT"]
        if len(recent) < 3:
            return None
        last3 = recent[-3:]
        if not (last3[0] == last3[1] == last3[2]):
            return None

        prev = data1["trendInfo"]["previousT"]
        current = data1["trendInfo"]["currentT"]
        price = float(data1["trendInfo"]["c"] or 0)
        planned = int(data1["common"]["plannedQuantity"] or 1)
        avail_qty = int(data1["common"]["availableQuantity"] or 0)

        # ATR value if present
        atr_col = f"ATR_10"
        atr_value = None
        if atr_col in trend_df.columns:
            try:
                atr_value = float(trend_df[atr_col].iloc[-1])
            except Exception:
                atr_value = None
        if atr_value is None and price > 0:
            # fallback small value
            atr_value = max(price * 0.002, 0.1)

        stop_dist = atr_value * atr_multiplier

        # Enter when flat and new confirmed trend (previousT different or empty)
        if avail_qty == 0 and (prev != last3[0] or prev == ""):
            if last3[0]:  # uptrend -> BUY
                stop = data1["trendInfo"].get("sl") or max(price - stop_dist, 0)
                return {"action": "BUY", "reverse": False, "qty": planned, "stop": stop, "reason": "enter_long_confirmed"}
            else:  # downtrend -> SELL
                stop = data1["trendInfo"].get("sl") or (price + stop_dist)
                return {"action": "SELL", "reverse": False, "qty": planned, "stop": stop, "reason": "enter_short_confirmed"}

        # If position exists and trend reversed -> consider conservative reverse
        if avail_qty != 0 and last3[0] != prev:
            # compute reverse qty conservatively
            reverse_qty = min(abs(avail_qty) * 2, planned * max_layers)
            reverse_qty = max(int(reverse_qty), planned)
            if last3[0]:
                stop = data1["trendInfo"].get("sl") or max(price - stop_dist, 0)
                return {"action": "BUY", "reverse": True, "qty": reverse_qty, "stop": stop, "reason": "reverse_to_long"}
            else:
                stop = data1["trendInfo"].get("sl") or (price + stop_dist)
                return {"action": "SELL", "reverse": True, "qty": reverse_qty, "stop": stop, "reason": "reverse_to_short"}

        return None

    def writeExcel(self, dataObj):
        current_date = datetime.datetime.now().strftime("%Y-%m-%d")
        stock = dataObj.get("stock", "unknown_stock")
        filename = f"data_{current_date}_{stock}_EQ2.xlsx"
        new_df = pd.DataFrame([dataObj])
        if not os.path.exists(filename):
            new_df.to_excel(filename, index=False)
        else:
            try:
                # Simple and portable: read existing, concat, overwrite file
                existing = pd.read_excel(filename)
                result = pd.concat([existing, new_df], ignore_index=True)
                result.to_excel(filename, index=False)
            except Exception:
                # Fallback: append to sheet without reading whole file (requires openpyxl)
                from openpyxl import load_workbook
                book = load_workbook(filename)
                with pd.ExcelWriter(filename, engine='openpyxl', mode='a') as writer:
                    # startrow: append after last used row
                    startrow = book.active.max_row
                    new_df.to_excel(writer, index=False, header=False, startrow=startrow)

@app.route('/')
def index():
    return render_template('Equity2.html')

@app.route('/getDataEQ2', methods=['POST'])
def mainMethod():            
    global enctoken, kite, kiteNew, payload
    payload = request.get_json(force=False, silent=True) 
    print('Received data:', json.dumps(payload))

    if payload is None:
        return json.dumps({"error": "invalid or missing JSON payload"})

    data1["stock"]["stockName"] = payload.get('stock')
    data1["common"]["plannedQuantity"] = int(payload.get('quantity'))
    data1["stock"]["superTrend"] = payload.get('sprTrend')
    data1["stock"]["stoploss"] = -abs(int(payload.get('stoploss')))

    equity = Equity()
    if enctoken == "" and payload.get('encToken'):
        equity.setEnctoken(payload.get('encToken'))
        print('SET ENCTOKEN ONE TIME')

    try:                        
        equity.setAllData() 
        data1["common"]["dateTime"] = datetime.datetime.now().strftime("%Y-%m-%d %I:%M:%S %p")
        data1["common"]["message"] = ""
        data1["common"]["transactionType"] = ""

        # Check Exit Conditions: Stop Loss OR Time-based Exit
        current_time = datetime.datetime.now().time()
        exit_time = datetime.time(15, 10)  # 3:10 PM
        
        stop_loss_triggered = data1["common"]["currentPnl"] < data1["stock"]["stoploss"]
        time_exit_triggered = current_time >= exit_time and not data1["StopAlgo"]
        
        if stop_loss_triggered or time_exit_triggered:
            data1["StopAlgo"] = True
            exit_reason = "Stop Loss Reached" if stop_loss_triggered else "3:10 PM - Market Exit Time"            
            if data1["common"]["availableQuantity"] != 0:
                if data1["common"]["availableQuantity"] > 0:
                    equity.placeOrder("SELL")
                    data1["common"]["transactionType"] = "SELL"
                    data1["common"]["message"] = f"{exit_reason} - Closing Long Position"
                else:
                    equity.placeOrder("BUY")
                    data1["common"]["transactionType"] = "BUY"
                    data1["common"]["message"] = f"{exit_reason} - Closing Short Position"

        # use decide_action to centralize decision logic
        if not data1["StopAlgo"]:
            decision = equity.decide_action()
        else:
            decision = None
        print('Decision:', decision)
        if decision:
            # place order according to decision
            if decision.get("reverse"):
                # additionalQty will add plannedQuantity (keeps existing placeOrder signature)  
                equity.placeOrder(decision["action"], additionalQty=data1["common"]["plannedQuantity"])
                print('')
            else:
                equity.placeOrder(decision["action"])
                print('')

            data1["trendInfo"]["previousT"] = data1["trendInfo"]["currentT"]
            data1["common"]["message"] = f"{decision['action']} Order Placed ({decision.get('reason')})"
            data1["common"]["transactionType"] = decision["action"]
        else:
            data1["common"]["message"] = "No action (no confirmed signal or insufficient state)"
            print("No Trend Change or insufficient confirmation")
    except Exception as inst:
        print(type(inst))    # the exception type
        print(inst.args)     # arguments stored in .args
        print(inst)          # __str__ allows args to be printed directly,
    finally:        
        print("------------------------------------------------------------")
        report = {"dateTime": data1["common"]["dateTime"],
                      "stock": data1["stock"]["stockName"],
                      "trend": data1["trendInfo"]["currentT"],
                      "open": data1["trendInfo"]["o"],
                      "high": data1["trendInfo"]["h"],
                      "low": data1["trendInfo"]["l"],
                      "close": data1["trendInfo"]["c"],
                      "average": data1["trendInfo"]["a"],
                      "stopLoss": data1["trendInfo"]["sl"],
                      "transactionType": data1["common"]["transactionType"],
                      "message": data1["common"]["message"],
                      "availableQty": data1["common"]["availableQuantity"],
                      "currentPnl": data1["common"]["currentPnl"],
                      "stopAlgo": data1["StopAlgo"]
                      }
        equity.writeExcel(report)
        return json.dumps(data1, default=str)

if __name__ == '__main__':    
    app.run(debug=True, port=5002)