import os
import datetime
import pandas as pd


def write_excel(dataObj, base_dir):
    """
    Appends dataObj as a row to the daily Excel report in the Reports subfolder.

    Args:
        dataObj: dict of report data to write.
        base_dir: directory of the calling script (__file__'s dirname).
    """
    current_date = datetime.datetime.now().strftime("%Y-%m-%d")
    stock = dataObj.get("stock", "unknown_stock")
    reports_dir = os.path.join(base_dir, "Reports")
    os.makedirs(reports_dir, exist_ok=True)
    filename = os.path.join(reports_dir, f"data_{current_date}_{stock}_EQ2.xlsx")
    new_df = pd.DataFrame([dataObj])
    if not os.path.exists(filename):
        new_df.to_excel(filename, index=False)
    else:
        try:
            existing = pd.read_excel(filename)
            result = pd.concat([existing, new_df], ignore_index=True)
            result.to_excel(filename, index=False)
        except Exception:
            from openpyxl import load_workbook
            book = load_workbook(filename)
            with pd.ExcelWriter(filename, engine='openpyxl', mode='a') as writer:
                startrow = book.active.max_row
                new_df.to_excel(writer, index=False, header=False, startrow=startrow)
